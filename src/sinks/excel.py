"""Excel (.xlsx) sink — senior ko'rinish: ko'k sarlavha, yashil/qizil qatorlar,
№ avtomatik raqamlash, toza formatlar. openpyxl bilan."""
from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from .base import Sink
from ..models import TRADE_SIGN_COL, OPEN_SIGN_COL

log = logging.getLogger(__name__)

# ── Rang palitrasi (senior, toza) ──
HEADER_FILL = PatternFill("solid", fgColor="4472C4")   # ko'k
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
WIN_FILL = PatternFill("solid", fgColor="A9D08E")      # yashil (yutuq)
LOSS_FILL = PatternFill("solid", fgColor="F4796B")     # qizil (zarar)
CENTER = Alignment(horizontal="center", vertical="center")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# Ustun kengliklari (sarlavha nomi -> kenglik)
_WIDTHS = {
    "№": 5, "SANA": 11, "LOT": 7, "ENTRY": 11, "OCHILISH VAQTI": 15,
    "TYPE": 8, "TP": 11, "SL": 11, "YOPILISH VAQTI": 15, "RESULT": 9,
    "PIPS": 8, "P/L $": 12, "HOZIRGI NARX": 13, "FLOATING $": 12,
}


class ExcelSink(Sink):
    name = "excel"

    def __init__(self, path: str | Path, trades_sheet: str = "Trades", open_sheet: str = "Open_Positions"):
        self.path = Path(path)
        self.trades_sheet = trades_sheet
        self.open_sheet = open_sheet
        self._trade_header: list[str] = []
        self._open_header: list[str] = []

    # ── tayyorlash ──
    def ensure_ready(self, trade_header: list[str], open_header: list[str]) -> None:
        self._trade_header = trade_header
        self._open_header = open_header
        self.path.parent.mkdir(parents=True, exist_ok=True)

        if not self.path.exists():
            wb = Workbook()
            ws = wb.active
            ws.title = self.trades_sheet
            self._style_header(ws, trade_header)
            self._style_header(wb.create_sheet(self.open_sheet), open_header)
            wb.save(self.path)
            log.info("Excel yaratildi: %s", self.path)
            return

        wb = load_workbook(self.path)
        for title, header in ((self.trades_sheet, trade_header), (self.open_sheet, open_header)):
            if title not in wb.sheetnames:
                self._style_header(wb.create_sheet(title), header)
        wb.save(self.path)

    def _style_header(self, ws, header: list[str]) -> None:
        ws.append(header)
        for idx, name in enumerate(header, start=1):
            cell = ws.cell(row=1, column=idx)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = CENTER
            cell.border = BORDER
            ws.column_dimensions[cell.column_letter].width = _WIDTHS.get(name, 12)
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 22

    # ── yopilgan savdolar ──
    def append_trades(self, rows: list[list[Any]]) -> None:
        if not rows:
            return
        wb = load_workbook(self.path)
        ws = wb[self.trades_sheet]
        sign_idx = self._trade_header.index(TRADE_SIGN_COL)  # RESULT ustuni (0-based, № bilan)
        start_no = ws.max_row  # header=1 => birinchi savdo № = 1
        for i, data_row in enumerate(rows):
            no = start_no + i
            full = [no, *data_row]
            ws.append(full)
            is_win = str(full[sign_idx]).strip().startswith("+")
            self._paint_row(ws, ws.max_row, len(full), is_win)
        wb.save(self.path)
        log.info("Excel: %d savdo qo'shildi", len(rows))

    def reset_trades(self) -> None:
        """Trades varag'ini tozalaydi (faqat sarlavha qoladi) — qayta tiklash uchun."""
        wb = load_workbook(self.path) if self.path.exists() else Workbook()
        if self.trades_sheet in wb.sheetnames:
            del wb[self.trades_sheet]
        ws = wb.create_sheet(self.trades_sheet, 0)
        self._style_header(ws, self._trade_header)
        wb.save(self.path)

    # ── ochiq pozitsiyalar (to'liq qayta yoziladi) ──
    def replace_open_positions(self, rows: list[list[Any]]) -> None:
        wb = load_workbook(self.path)
        if self.open_sheet in wb.sheetnames:
            del wb[self.open_sheet]
        ws = wb.create_sheet(self.open_sheet)
        self._style_header(ws, self._open_header)
        sign_idx = self._open_header.index(OPEN_SIGN_COL)  # FLOATING $ ustuni
        for i, data_row in enumerate(rows):
            full = [i + 1, *data_row]
            ws.append(full)
            is_win = not str(full[sign_idx]).strip().startswith("-")
            self._paint_row(ws, ws.max_row, len(full), is_win)
        wb.save(self.path)

    # ── qatorni bo'yash + format ──
    def _paint_row(self, ws, row_no: int, ncols: int, is_win: bool) -> None:
        fill = WIN_FILL if is_win else LOSS_FILL
        for col in range(1, ncols + 1):
            cell = ws.cell(row=row_no, column=col)
            cell.fill = fill
            cell.border = BORDER
            cell.alignment = CENTER
            name = self._trade_header[col - 1] if col <= len(self._trade_header) else ""
            if name in ("LOT",):
                cell.number_format = "0.00"
            elif name in ("№", "PIPS"):
                cell.number_format = "0"
